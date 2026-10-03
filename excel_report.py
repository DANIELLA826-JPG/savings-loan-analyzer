"""
excel_report.py
Step 6: Combine the results into one formatted Excel report.
Run the earlier scripts first so the CSV files and charts exist.
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill, Alignment

OUTPUT_FILE = "savings_report.xlsx"

# 1. Load the results from the earlier steps
monthly = pd.read_csv("monthly_summary.csv")
scenarios = pd.read_csv("loan_scenarios.csv")

data = pd.read_csv("savings_data.csv", parse_dates=["date"])
data["weekday"] = data["date"].dt.day_name()
order = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
by_weekday = (
    data.groupby("weekday")["amount"].sum().round(2).reindex(order).reset_index()
)
by_weekday.columns = ["weekday", "total_saved"]

# 2. Write each table to its own sheet
with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
    monthly.to_excel(writer, sheet_name="Monthly Summary", index=False)
    by_weekday.to_excel(writer, sheet_name="By Weekday", index=False)
    scenarios.to_excel(writer, sheet_name="Loan Scenarios", index=False)

# 3. Reopen the file to make it look professional
wb = load_workbook(OUTPUT_FILE)
header_fill = PatternFill("solid", start_color="1F3864")
header_font = Font(bold=True, color="FFFFFF")

for sheet in wb.worksheets:
    # Style the header row
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    # Make every column wide enough for its contents
    for column in sheet.columns:
        longest = max(len(str(cell.value)) if cell.value is not None else 0
                      for cell in column)
        sheet.column_dimensions[column[0].column_letter].width = max(longest + 4, 14)

    # Keep the header visible while scrolling
    sheet.freeze_panes = "A2"

# 4. Add a sheet with the chart images
charts = wb.create_sheet("Charts")
charts["A1"] = "Savings & Loan Rate Analysis (simulated data)"
charts["A1"].font = Font(bold=True, size=14)

for position, file in [("A3", "charts/monthly_savings.png"),
                       ("A30", "charts/savings_by_weekday.png"),
                       ("A57", "charts/loan_scenarios.png")]:
    picture = Image(file)
    picture.width, picture.height = 600, 333   # shrink to fit the sheet
    charts.add_image(picture, position)

wb.save(OUTPUT_FILE)
print(f"Created {OUTPUT_FILE} with 4 sheets.")
